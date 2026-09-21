import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_border(cell):
    """Sets clean thin border for table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        f'<w:left w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

def create_document():
    doc = Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base Normal Style: 100% Black font, Calibri
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x00, 0x00, 0x00) # Pure Black

    # -------------------------------------------------------------
    # Document Header / Title (Pure Black & White)
    # -------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("Project 2: Distributed Microservices Deployment on AWS EC2")
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(14)
    run_sub = sub_p.add_run("Continuation from Project 1: 5 Microservices across 3 Ubuntu EC2 Instances with API Gateway & Role-Based JWT Authentication")
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    # Student Info Table
    info_table = doc.add_table(rows=4, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_data = [
        ("Course / Project:", "Cloud Computing / Distributed Microservices Project 2"),
        ("Student Name:", "Sokhadom Khorn"),
        ("Student ID:", "2023517"),
        ("Date:", "September 2026")
    ]
    for idx, (label, val) in enumerate(info_data):
        row = info_table.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        set_cell_border(c0)
        set_cell_border(c1)
        c0.text = label
        c0.paragraphs[0].runs[0].font.bold = True
        c0.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        c1.text = val
        c1.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        set_cell_background(c0, "F5F5F5")
        set_cell_background(c1, "FFFFFF")
    
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # 1. Multi-Instance Architecture Table
    # -------------------------------------------------------------
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    r_h1 = h1.add_run("1. Multi-Instance Architecture & IP Mapping")
    r_h1.font.size = Pt(14)
    r_h1.font.bold = True
    r_h1.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    table = doc.add_table(rows=4, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["EC2 Instance", "Operating System", "Hosted Microservices", "Public IPv4 Address"]
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        set_cell_border(hdr_cells[i])
        hdr_cells[i].text = title
        p_hdr = hdr_cells[i].paragraphs[0]
        p_hdr.runs[0].font.bold = True
        p_hdr.runs[0].font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        set_cell_background(hdr_cells[i], "EBEBEB")

    rows_data = [
        ("EC2 Instance 1 (Gateway/Reg/Login)", "Ubuntu 22.04 LTS", "API Gateway (:4000)\nRegistration (:5001)\nLogin (:5002)\nMongoDB (:27017)", "Public: 44.214.139.176\n(Private: 172.31.9.169)"),
        ("EC2 Instance 2 (Admin Service)", "Ubuntu 22.04 LTS", "Admin Microservice (:5003)", "Public: 100.57.160.206\n(Private: 172.31.3.233)"),
        ("EC2 Instance 3 (User Service)", "Ubuntu 22.04 LTS", "User Microservice (:5004)", "Public: 44.192.127.185\n(Private: 172.31.2.148)")
    ]

    for row_idx, rdata in enumerate(rows_data):
        row_cells = table.rows[row_idx + 1].cells
        for col_idx, text in enumerate(rdata):
            set_cell_border(row_cells[col_idx])
            row_cells[col_idx].text = text
            row_cells[col_idx].paragraphs[0].runs[0].font.color.rgb = RGBColor(0x00, 0x00, 0x00)
            if row_idx % 2 == 1:
                set_cell_background(row_cells[col_idx], "FAFAFA")
            else:
                set_cell_background(row_cells[col_idx], "FFFFFF")

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # -------------------------------------------------------------
    # Helper to add section with screenshot (Black & White styling)
    # -------------------------------------------------------------
    screenshots_dir = os.path.join(os.path.dirname(__file__), "screenshots")

    def add_task_section(title, task_num, description, image_filename, expected_output=None):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        r_task = h.add_run(f"{task_num}: {title}")
        r_task.font.size = Pt(13)
        r_task.font.bold = True
        r_task.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

        desc_p = doc.add_paragraph()
        desc_p.paragraph_format.space_after = Pt(4)
        r_desc = desc_p.add_run(description)
        r_desc.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

        if expected_output:
            out_p = doc.add_paragraph()
            out_p.paragraph_format.space_after = Pt(6)
            out_run = out_p.add_run(f"Expected Output / Status: {expected_output}")
            out_run.font.bold = True
            out_run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

        image_path = os.path.join(screenshots_dir, image_filename)
        if os.path.exists(image_path):
            cap_p = doc.add_paragraph()
            cap_p.paragraph_format.space_after = Pt(2)
            r_cap = cap_p.add_run(f"Figure: {title} Screenshot")
            r_cap.font.italic = True
            r_cap.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
            doc.add_picture(image_path, width=Inches(6.2))
        else:
            pbox = doc.add_table(rows=1, cols=1)
            pbox.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = pbox.rows[0].cells[0]
            set_cell_border(cell)
            set_cell_background(cell, "F5F5F5")
            cp = cell.paragraphs[0]
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            c_run = cp.add_run(f"\n[ SCREENSHOT PLACEHOLDER: {title} ]\n")
            c_run.font.italic = True
            c_run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        
        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # -------------------------------------------------------------
    # Task Sections matching user instructions exactly
    # -------------------------------------------------------------

    # Task 1
    h_t1 = doc.add_paragraph()
    h_t1.paragraph_format.space_before = Pt(14)
    h_t1.paragraph_format.space_after = Pt(4)
    r_t1 = h_t1.add_run("Task 1: Ubuntu OS Configuration on all 3 EC2 Instances")
    r_t1.font.size = Pt(13)
    r_t1.font.bold = True
    r_t1.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    p_t1 = doc.add_paragraph()
    p_t1.paragraph_format.space_after = Pt(8)
    r_p_t1 = p_t1.add_run(
        "All 3 EC2 instances were launched with Ubuntu 22.04 LTS OS and fully configured using automated scripts:\n"
        "1. Ubuntu package manager update and upgrade: 'sudo apt-get update && sudo apt-get upgrade -y'\n"
        "2. Node.js (v20 LTS) and npm installation via official NodeSource repository\n"
        "3. Verification of Git installation ('git --version' -> git version 2.34.1)\n"
        "4. Process management with PM2 and MongoDB Community Edition on port 27017 (bind: 0.0.0.0)"
    )
    r_p_t1.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    # Task 2
    add_task_section(
        "EC2 Instance 1 Public IP",
        "Task 2.1",
        "AWS EC2 Management Console screenshot showing EC2 Instance 1 (API Gateway, Registration, Login) running with Public IPv4 address: 44.214.139.176.",
        "task2_ec2_1_ip.png"
    )

    add_task_section(
        "EC2 Instance 2 Public IP",
        "Task 2.2",
        "AWS EC2 Management Console screenshot showing EC2 Instance 2 (Admin Microservice) running with Public IPv4 address: 100.57.160.206.",
        "task2_ec2_2_ip.png"
    )

    add_task_section(
        "EC2 Instance 3 Public IP",
        "Task 2.3",
        "AWS EC2 Management Console screenshot showing EC2 Instance 3 (User Microservice) running with Public IPv4 address: 44.192.127.185.",
        "task2_ec2_3_ip.png"
    )

    # Task 3 & 4
    h_t4 = doc.add_paragraph()
    h_t4.paragraph_format.space_before = Pt(14)
    h_t4.paragraph_format.space_after = Pt(4)
    r_t4 = h_t4.add_run("Task 4: Updated GitHub Repository & Code Deployment")
    r_t4.font.size = Pt(13)
    r_t4.font.bold = True
    r_t4.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    p_t4 = doc.add_paragraph()
    p_t4.paragraph_format.space_after = Pt(8)
    r_p_t4 = p_t4.add_run(
        "The project source code in GitHub was updated with the downstream target EC2 Public IPs configured in 'APIGateway_Microservice/index.js'.\n"
        "GitHub Repository URL: https://github.com/SokhadomKhorn/University_Project2\n"
        "Commands executed on EC2 instances:\n"
        "  git clone https://github.com/SokhadomKhorn/University_Project2.git\n"
        "  cd University_Project2\n"
        "  npm run install:all"
    )
    r_p_t4.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    # Task 5.1
    add_task_section(
        "Add a New Record (POST /register/userregister)",
        "Task 5.1",
        "Postman executing 'POST http://44.214.139.176:4000/register/userregister' registering a new user record through the API Gateway.",
        "task5_1_add_record_postman.png",
        "HTTP 201 Created - User registered successfully"
    )

    add_task_section(
        "MongoDB Document Verification for New Record",
        "Task 5.1 (MongoDB)",
        "Terminal executing 'mongosh university_identity_db' displaying the registered user documents with salted bcrypt password hashes and roles.",
        "task5_1_mongodb.png"
    )

    # Task 5.2
    add_task_section(
        "Duplicate Email Registration Rejection",
        "Task 5.2",
        "Postman executing 'POST http://44.214.139.176:4000/register/userregister' with the duplicate email ID. Verifies uniqueness constraint.",
        "task5_2_duplicate_email.png",
        "HTTP 400 Bad Request - Duplicate Email: An account with this email address already exists."
    )

    # Task 6.1
    add_task_section(
        "Login as Admin and Generate JWT",
        "Task 6.1",
        "Postman executing 'POST http://44.214.139.176:4000/auth/login' with Admin credentials and role 'admin'. Returns signed JWT.",
        "task6_1_admin_login_jwt.png",
        "HTTP 200 OK - Login successful with Admin JWT Token"
    )

    # Task 6.2
    add_task_section(
        "Login as User and Generate JWT",
        "Task 6.2",
        "Postman executing 'POST http://44.214.139.176:4000/auth/login' with User credentials and role 'user'. Returns signed JWT.",
        "task6_2_user_login_jwt.png",
        "HTTP 200 OK - Login successful with User JWT Token"
    )

    # Task 7.1
    add_task_section(
        "Admin + Admin JWT call viewalluser API",
        "Task 7.1",
        "Postman executing 'GET http://44.214.139.176:4000/admin/viewalluser' routed by API Gateway to EC2 Instance 2 (100.57.160.206:5003) with Admin JWT.",
        "task7_1_admin_viewalluser.png",
        "HTTP 200 OK - All users retrieved successfully"
    )

    # Task 7.2
    add_task_section(
        "User + User JWT call updateprofile API",
        "Task 7.2",
        "Postman executing 'PUT http://44.214.139.176:4000/user/updateprofile' routed by API Gateway to EC2 Instance 3 (44.192.127.185:5004) with User JWT.",
        "task7_2_user_updateprofile.png",
        "HTTP 200 OK - User profile updated successfully"
    )

    # Task 8.1
    add_task_section(
        "Admin calling User updateprofile API (RBAC Verification)",
        "Task 8.1",
        "Postman executing 'PUT http://44.214.139.176:4000/user/updateprofile' using Admin JWT. API Gateway enforces RBAC and rejects request.",
        "task8_1_admin_call_user_api.png",
        "HTTP 403 Forbidden - Access Denied: Role 'user' is required to access this endpoint. Your role is 'admin'."
    )

    # Task 8.2
    add_task_section(
        "User calling Admin viewalluser API (RBAC Verification)",
        "Task 8.2",
        "Postman executing 'GET http://44.214.139.176:4000/admin/viewalluser' using User JWT. API Gateway enforces RBAC and rejects request.",
        "task8_2_user_call_admin_api.png",
        "HTTP 403 Forbidden - Access Denied: Role 'admin' is required to access this endpoint. Your role is 'user'."
    )

    output_doc_path = os.path.join(os.path.dirname(__file__), "Project2_Submission.docx")
    doc.save(output_doc_path)
    print(f"[SUCCESS] Word document generated at: {output_doc_path}")

    # Convert to PDF
    try:
        from docx2pdf import convert
        output_pdf_path = os.path.join(os.path.dirname(__file__), "Project2_Submission.pdf")
        convert(output_doc_path, output_pdf_path)
        print(f"[SUCCESS] PDF document generated at: {output_pdf_path}")
    except Exception as e:
        print(f"[NOTE] Automated docx2pdf conversion skipped: {e}")

if __name__ == '__main__':
    create_document()
